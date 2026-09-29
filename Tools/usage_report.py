"""THE USAGE SHEET (the CEO's ask 2026-09-03: one aggregate sheet - totals
per day, week, and month for every model and every tool, CSV for humans +
TXT twin, scripted).

WHERE THE NUMBERS COME FROM: the Claude Code harness writes a full JSONL
transcript of every session under ~/.claude/projects/<project>/*.jsonl.
Every assistant message in there carries the REAL API meter (model, input /
output / thinking tokens, cache reads and writes) and every tool call is
named - including sub-agent (employee) traffic, which is flagged and metered
under its own model. So unlike SUBAGENTS.md's self-estimates, everything in
this sheet is harness-metered truth for THIS machine. (Each workstation only
has its own transcripts; the sheet keeps WS1 and WS2 rows side by side and a
run only rewrites its own workstation's rows.)

WHAT A "TOOL COST" MEANS: a tool call itself is free; its cost is the result
text injected into context (estimated at chars/4), which is then re-read by
every later turn via cache. calls + context_est_tok is therefore the honest
per-tool figure; there is no per-tool dollar meter.

COSTS ARE ESTIMATES: token counts are real; the est_cost_usd column prices
them from the PRICING table below (blank when a model's price is not filled
in). The billing dashboard remains the only truth for dollars.

Outputs (REGENERATED whole each run - the transcripts are the ledger, this
is the derived sheet, so no append-only file here):
  docs/history/usage_metrics.txt  - OPEN THIS ONE: totals by day/week/month,
                                    then the BREAKDOWNS (the CEO's ask
                                    2026-09-10): averages per request, the
                                    context window by month, averages per
                                    tool call, employee runs
  docs/history/usage_metrics.xlsx - THE SPREADSHEET (the CEO's ask
                                    2026-09-10): header row frozen, every
                                    token column a Number with thousands
                                    separators, per-period TOTAL rows bold
                                    with a thick bottom border, total_tok +
                                    avg_tok per model/tool in columns O/P;
                                    sheets usage / per_request / employees
  docs/history/usage_metrics.csv  - the same rows in long format (+ avg_req
                                    rows per month/all)
  docs/history/usage_employees.csv - one row per sub-agent run (day, model,
                                    minutes, tokens, tool calls, est$, the
                                    brief's first line)
  docs/history/usage_daily.txt    - THE DAILY LINE (the CEO 2026-09-10, ideas
                                    17 + 20): one line per active day - the
                                    weighted total, its top pillar, cache
                                    misses, the read diet - each judged
                                    against the previous 7 active days
                                    (HIGH / normal / LOW). standup prints its
                                    tail; "a number without comparison means
                                    nothing".
  docs/history/usage_cache.json   - incremental parse cache (transcripts are
                                    append-only; reruns only read new bytes)
  docs/history/usage_by_arc.txt   - THE PER-ARC LINE (the CEO 2026-09-14):
                                    one line per arc between consecutive
                                    "Checkpoint:" commits - commits, weighted
                                    spend, employee runs, spend per commit,
                                    top pillar, a CHECK verdict vs the
                                    previous 5 arcs; the open arc since the
                                    last checkpoint is the final "(open)" line
  docs/history/cache_misses.txt   - THE CACHE-MISS INVESTIGATION (the CEO
                                    2026-09-14): one line per detected cache
                                    miss, classified by likely cause (session
                                    start, TTL gap, long pause, prefix change
                                    mid-session, employee) with a SUMMARY
                                    block; regenerated whole each run
  docs/history/cache_miss_runs.txt - append-only, one line per run: misses,
                                    wasted ~weighted, dominant cause

THE WEIGHTED COLUMN (the CEO's insight 2026-09-10, idea 13): only the tokens
that count against the plan matter. weighted_tok prices every token relative
to fresh input: input 1, cache write 1.25 (2.0 on the 1-hour TTL), cache
read 0.1 (0.025 on Fable 5.1), output 5. It is THE headline; the raw count
is trivia beside it. The same weights drive the fan-out guard.

CACHE MISSES (idea 14): a request whose cache WRITE is most of its box is a
prefix rewrite - the cache expired, or an early byte of the prefix (the
core instructions file, a memory file) changed mid-session. The first
request of a session is a COLD start (expected, counted apart); every later
one is a MISS. The sheet counts both and prices the miss tokens, because
cache writes are the biggest weighted pillar and misses are the only waste
nobody chose.

THE READ DIET (idea 17): every Read call carries its parameters, so the
sheet grades the wiki convention - section reads (offset/limit, sed -n,
head/tail) vs whole-file reads (bare Read, cat, type, Get-Content) - and the
context each injected. Shell reads count too.

WHAT "THE CONTEXT WINDOW" MEANS HERE: the harness never records what sits
in the window, but every request's prompt size is input + cache_read +
cache_create - that IS the box the model read that turn. The sheet reports
its average and its biggest per model per month; a growing average is a
session that should have checkpointed.

Usage:  python tools/usage_report.py            # in run_all's metrics group
Needs:  openpyxl (pip) for the .xlsx; without it the CSV/TXT still write
        and the run says so (WORKSTATION.md carries the row).

Usage:  python tools/usage_report.py --quiet   # standup's silent refresh

PURPOSE: Mines the Claude Code harness JSONL transcripts for real per model
  per tool token usage and writes an aggregate usage sheet (CSV, TXT, XLSX,
  employee runs, and the daily budget line) totaled by day, week and month,
  plus usage_by_arc.txt (checkpoint to checkpoint) and cache_misses.txt +
  cache_miss_runs.txt (every miss with its likely cause).
INTENT: the CEO's ask 2026-09-03: 'one aggregate sheet, totals per day,
  week, and month for every model and every tool, CSV for humans plus TXT
  twin, scripted.'

Search keys: usage, token costs, tool costs, metrics sheet, daily totals,
weekly totals, monthly totals, harness meter, transcript mining, average
tokens per request, per tool call, per employee, context window size,
weighted tokens, budget, cache misses, read diet, daily line, comparison.
See also: tools/metrics_report.py (task outcomes from SUBAGENTS.md);
REPORTING_METHOD.md; TOKEN_IDEAS.md.
"""
import csv
import datetime
import json
import os
import re
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HIST = os.path.join(ROOT, "docs", "history")
CSV_OUT = os.path.join(HIST, "usage_metrics.csv")
TXT_OUT = os.path.join(HIST, "usage_metrics.txt")
EMP_OUT = os.path.join(HIST, "usage_employees.csv")
XLSX_OUT = os.path.join(HIST, "usage_metrics.xlsx")
CACHE = os.path.join(HIST, "usage_cache.json")
DAILY_OUT = os.path.join(HIST, "usage_daily.txt")
ARC_OUT = os.path.join(HIST, "usage_by_arc.txt")
MISS_OUT = os.path.join(HIST, "cache_misses.txt")
MISS_RUNS_OUT = os.path.join(HIST, "cache_miss_runs.txt")
CACHE_VERSION = 5  # v3: weighted, cache misses, read-diet classes; v4: images
# priced by pixels; v5: per-request timestamp log (reqs) for the per-arc
# ledger + per-miss detail records (misses) for the cache-miss investigation
BASELINE_ARCS = 5      # per-arc verdict baseline window
ARC_OVERSPEND = 1.50   # per-commit spend past this x the baseline mean is CHECK

# $ per MILLION tokens: model -> (input, output). None = unknown; fill from
# the billing page / claude.com/pricing and the cost column comes alive.
# Cache pricing law: read = 0.1x input (0.025x on Fable 5.1); 5-min write =
# 1.25x; 1-hour write = 2x.
# Rates confirmed 2026-09-04 against Anthropic's published first-party API
# pricing (haiku 4.5 / sonnet 5 / opus 5 / fable 5). The old opus row was
# the Opus 4.1-era 15/75 - Opus 5-tier is 5/25.
# Third value = the CACHE READ multiplier on the input price: 0.1x on most
# models, 0.025x on Claude Fable 5.1 ($0.25/MTok - every "fable" message in
# these transcripts is claude-fable-5-1; corrected 2026-09-10 after the CEO
# asked whether cache reads count: they do, at this discount).
PRICING = {
    "haiku": (1.0, 5.0, 0.1),
    "sonnet": (3.0, 15.0, 0.1),
    "opus": (5.0, 25.0, 0.1),
    "fable": (10.0, 50.0, 0.025),
}

# THE WEIGHTS (idea 13): what a token costs relative to fresh input. Cache
# reads take the model's multiplier from PRICING (third value).
W_INPUT, W_WRITE_5M, W_WRITE_1H, W_OUTPUT = 1.0, 1.25, 2.0, 5.0
DEFAULT_READ_MULT = 0.1
MISS_SHARE = 0.5       # a request whose cache write is >= this share of its
MISS_MIN_BOX = 20000   # box (and the box is at least this big) is a MISS
BIG_READ_TOK = 10000   # idea 15: a whole-file read past this is "heavy"
COMPARE_DAYS = 7       # the daily line compares against this many prior active days
HIGH_RATIO, LOW_RATIO = 1.30, 0.70

FIELDS = ["period_type", "period", "ws", "scope", "name", "role", "count",
          "input_tok", "output_tok", "think_tok", "cache_read_tok",
          "cache_create_tok", "context_est_tok", "est_cost_usd",
          "total_tok", "avg_tok", "weighted_tok", "cache_miss_n",
          "cache_miss_tok"]
# total_tok: a model row = every token metered (in + out + cache_rd +
# cache_wr); a tool row = the context its results injected; a TOTAL row =
# the period's model tokens. avg_tok = total_tok per request / per call.
# weighted_tok = the budget-weighted count (THE headline). cache_miss_n /
# cache_miss_tok = prefix rewrites after the session's first request and
# the cache-write tokens they cost.
EMP_FIELDS = ["ws", "day", "started", "minutes", "parent_session", "model",
              "msgs", "input_tok", "output_tok", "think_tok",
              "cache_read_tok", "cache_create_tok", "tool_calls",
              "ctx_est_tok", "avg_box_tok", "est_cost_usd", "brief"]


def workstation():
    home = os.path.expanduser("~").lower()
    return "WS2" if "travis" in home else ("WS1" if "owner" in home else "WS?")


def transcript_dirs():
    base = os.path.join(os.path.expanduser("~"), ".claude", "projects")
    if not os.path.isdir(base):
        return []
    return [os.path.join(base, d) for d in os.listdir(base)
            if "everwood" in d.lower()
            and os.path.isdir(os.path.join(base, d))]


def short_model(model):
    if model and model.startswith("claude-"):
        return model.split("-")[1]
    return model or "?"


def local_day(ts):
    try:
        dt = datetime.datetime.fromisoformat(ts.replace("Z", "+00:00"))
        return dt.astimezone().strftime("%Y-%m-%d")
    except (ValueError, AttributeError):
        return None


def blank_model():
    return {"count": 0, "in": 0, "out": 0, "think": 0, "read": 0,
            "c5": 0, "c1": 0, "maxp": 0, "miss_n": 0, "miss_tok": 0,
            "cold_n": 0, "cold_tok": 0}


def read_mult(model):
    p = PRICING.get(model)
    return p[2] if p else DEFAULT_READ_MULT


def weighted(model, m):
    """Budget-weighted tokens for a bucket (idea 13)."""
    return (m["in"] * W_INPUT + m["out"] * W_OUTPUT
            + m["read"] * read_mult(model)
            + m["c5"] * W_WRITE_5M + m["c1"] * W_WRITE_1H)


def blank_reads():
    return {"whole": {"n": 0, "ctx": 0, "big": 0},
            "section": {"n": 0, "ctx": 0}, "grep": 0}


# The read-diet classifier (idea 17). Read: sectioned iff offset/limit.
# Shell: a bare cat/type/Get-Content is a whole-file read; sed -n, head,
# tail, -TotalCount/-Tail/Select -First are section reads.
_SHELL_READ = re.compile(r"(^|[\s;|&(])(cat|type|get-content|gc)\s+\S")
_SHELL_SECTION = re.compile(r"(^|[\s;|&(])(sed\s+-n|head\b|tail\b|awk\b)|-totalcount|-tail\b|select-object\s+-(first|last)|select\s+-(first|last)")


IMAGE_EXT = (".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp")


def image_tokens(path):
    """An image read is priced by PIXELS (~w*h/750 after the API downscales
    to a 1568 px long edge / ~1.15 MP), roughly 1-2k tokens for a screenshot.
    Sizing it by its base64 payload / 4 overstated it 50-100x and made every
    screenshot a "big whole-file read" (found 2026-09-10 by tools/big_reads.py:
    the 21-a-day "big reads" were PNG montages). A gone file -> the 1600 cap."""
    try:
        from PIL import Image
        with Image.open(path) as im:
            w, h = im.size
    except Exception:  # missing file, no PIL, unreadable image
        return 1600
    scale = min(1.0, 1568.0 / max(w, h), (1150000.0 / float(w * h)) ** 0.5)
    return max(1, int(w * scale * h * scale / 750.0))


def classify_read(tool, tin):
    """-> 'whole' | 'section' | 'grep' | None for one tool_use block."""
    tin = tin or {}
    if tool == "Read":
        return "section" if (tin.get("offset") or tin.get("limit")) else "whole"
    if tool == "Grep":
        return "grep"
    if tool in ("Bash", "PowerShell"):
        cmd = (tin.get("command") or "").lower()
        if _SHELL_READ.search(cmd):
            return "section" if _SHELL_SECTION.search(cmd) else "whole"
        if _SHELL_SECTION.search(cmd):
            return "section"
    return None


def box(m):
    """The prompt size the model read: fresh input + cache read + cache
    written. Summed over a bucket; divide by count for the average box."""
    return m["in"] + m["read"] + m["c5"] + m["c1"]


def first_text(msg, limit=90):
    """The first line of a user message (an employee's brief), for labels."""
    content = msg.get("content")
    text = ""
    if isinstance(content, str):
        text = content
    elif isinstance(content, list):
        for blk in content:
            if isinstance(blk, dict) and blk.get("type") == "text":
                text = blk.get("text") or ""
                break
    text = " ".join(text.split())
    return text[:limit] if text else ""


def parse_file(path, state, force_agent=False):
    """Stream one transcript, resuming from the cached byte offset.
    state = {"size", "offset", "last_id", "pending", "days", "reqs",
    "misses"} (mutated). force_agent: employee transcripts
    (subagents/agent-*.jsonl) meter under role "agent" even where records
    lack the isSidechain flag. "reqs" is a per-request timestamp log (one
    entry per metered assistant message: [ts, model, role, in, out, read,
    c5, c1]) that feeds THE PER-ARC LINE (usage_by_arc.txt), which needs
    finer-than-a-day time buckets. "misses" is one entry per detected cache
    miss ([ts, session8, gap_min, box, wr, wasted_weighted, cause]) that
    feeds cache_misses.txt."""
    size = os.path.getsize(path)
    if size == state.get("size") and state.get("days") is not None:
        return  # unchanged since last run
    if size < state.get("size", 0) or state.get("days") is None:
        state.update({"offset": 0, "last_id": "", "pending": {}, "days": {},
                      "t_first": "", "t_last": "", "label": "", "nreq": 0,
                      "prev_ts": "", "reqs": [], "misses": []})
    days, pending = state["days"], state["pending"]
    with open(path, "rb") as fh:
        fh.seek(state.get("offset", 0))
        for raw in fh:
            try:
                rec = json.loads(raw.decode("utf-8", "replace"))
            except (ValueError, UnicodeDecodeError):
                continue
            rtype = rec.get("type")
            ts = rec.get("timestamp") or ""
            if ts and rtype in ("assistant", "user"):
                if not state.get("t_first") or ts < state["t_first"]:
                    state["t_first"] = ts
                if ts > state.get("t_last", ""):
                    state["t_last"] = ts
            if rtype == "user" and not state.get("label"):
                state["label"] = first_text(rec.get("message") or {})
            if rtype == "assistant":
                msg = rec.get("message") or {}
                day = local_day(rec.get("timestamp", ""))
                if not day:
                    continue
                d = days.setdefault(day, {"models": {}, "tools": {}, "reads": blank_reads()})
                d.setdefault("reads", blank_reads())
                for blk in msg.get("content") or []:
                    if isinstance(blk, dict) and blk.get("type") == "tool_use":
                        tool = blk.get("name") or "?"
                        cls = classify_read(tool, blk.get("input"))
                        tin = blk.get("input") if isinstance(blk.get("input"), dict) else {}
                        pending[blk.get("id") or ""] = [tool, cls, tin.get("file_path")]
                        t = d["tools"].setdefault(tool, {"calls": 0, "ctx": 0})
                        t["calls"] += 1
                        if cls == "grep":
                            d["reads"]["grep"] += 1
                        elif cls:
                            d["reads"][cls]["n"] += 1
                # A message split across records repeats its meter; count once.
                mid = msg.get("id") or rec.get("requestId") or ""
                usage = msg.get("usage")
                if not usage or mid == state.get("last_id"):
                    continue
                state["last_id"] = mid
                model = short_model(msg.get("model"))
                if model in ("<synthetic>", "?"):
                    continue
                role = "agent" if (force_agent or rec.get("isSidechain")) \
                    else "main"
                m = d["models"].setdefault(model + "|" + role, blank_model())
                m["count"] += 1
                m["in"] += usage.get("input_tokens", 0) or 0
                m["out"] += usage.get("output_tokens", 0) or 0
                m["think"] += (usage.get("output_tokens_details") or {}
                               ).get("thinking_tokens", 0) or 0
                m["read"] += usage.get("cache_read_input_tokens", 0) or 0
                cc = usage.get("cache_creation") or {}
                c5 = cc.get("ephemeral_5m_input_tokens")
                c1 = cc.get("ephemeral_1h_input_tokens")
                if c5 is None and c1 is None:
                    c5 = usage.get("cache_creation_input_tokens", 0) or 0
                m["c5"] += c5 or 0
                m["c1"] += c1 or 0
                bx = ((usage.get("input_tokens", 0) or 0)
                      + (usage.get("cache_read_input_tokens", 0) or 0)
                      + (c5 or 0) + (c1 or 0))
                m["maxp"] = max(m["maxp"], bx)
                # Per-request timestamp log (v5): the per-arc ledger needs
                # finer-than-a-day time buckets (checkpoints fall mid-day).
                state.setdefault("reqs", []).append(
                    [ts, model, role, usage.get("input_tokens", 0) or 0,
                     usage.get("output_tokens", 0) or 0,
                     usage.get("cache_read_input_tokens", 0) or 0,
                     c5 or 0, c1 or 0])
                # Cache misses (idea 14): a big write on a not-first request.
                wr = (c5 or 0) + (c1 or 0)
                gap_min = minutes_between(state["prev_ts"], ts) \
                    if state.get("prev_ts") else None
                state["nreq"] = state.get("nreq", 0) + 1
                if bx >= MISS_MIN_BOX and wr >= MISS_SHARE * bx:
                    if state["nreq"] == 1:
                        m["cold_n"] += 1
                        m["cold_tok"] += wr
                        cause = "session start"
                    else:
                        m["miss_n"] += 1
                        m["miss_tok"] += wr
                        # Priority: an employee transcript's misses are
                        # expected fresh-context overhead, not investigable
                        # by pause length, so that check outranks the gap
                        # buckets even though it reads last in prose.
                        if force_agent:
                            cause = "employee"
                        elif gap_min is None or gap_min >= 60:
                            cause = "TTL gap"
                        elif gap_min >= 5:
                            cause = "long pause"
                        else:
                            cause = "prefix change mid-session"
                    wasted = (c5 or 0) * W_WRITE_5M + (c1 or 0) * W_WRITE_1H
                    state.setdefault("misses", []).append(
                        [ts, (rec.get("sessionId") or "?")[:8], gap_min,
                         bx, wr, wasted, cause])
                state["prev_ts"] = ts
            elif rtype == "user":
                # Tool results: size them - that is what enters context.
                msg = rec.get("message") or {}
                content = msg.get("content")
                if not isinstance(content, list):
                    continue
                day = local_day(rec.get("timestamp", ""))
                if not day:
                    continue
                for blk in content:
                    if not (isinstance(blk, dict)
                            and blk.get("type") == "tool_result"):
                        continue
                    ent = pending.pop(blk.get("tool_use_id") or "", None)
                    if not ent:
                        continue
                    ent = ent if isinstance(ent, list) else [ent]
                    tool, cls, fpath = (ent + [None, None, None])[:3]
                    payload = rec.get("toolUseResult")
                    if payload is None:
                        payload = blk.get("content")
                    if fpath and str(fpath).lower().endswith(IMAGE_EXT):
                        est = image_tokens(fpath)  # pixels, not base64 bytes
                    else:
                        try:
                            est = len(json.dumps(payload, default=str)) // 4
                        except (TypeError, ValueError):
                            est = len(str(payload)) // 4
                    d = days.setdefault(day, {"models": {}, "tools": {}, "reads": blank_reads()})
                    d.setdefault("reads", blank_reads())
                    t = d["tools"].setdefault(tool, {"calls": 0, "ctx": 0})
                    t["ctx"] += est
                    if cls in ("whole", "section"):
                        d["reads"][cls]["ctx"] += est
                        if cls == "whole" and est >= BIG_READ_TOK:
                            d["reads"]["whole"]["big"] += 1
        state["offset"] = fh.tell()
    state["size"] = size
    if len(pending) > 400:  # orphaned tool_use ids (interrupted calls)
        for k in list(pending)[:-200]:
            del pending[k]


def merge_bucket(tgt, d):
    for key, m in d["models"].items():
        t = tgt["models"].setdefault(key, blank_model())
        for k in m:
            if k == "maxp":
                t[k] = max(t[k], m[k])
            else:
                t[k] += m[k]
    for tool, v in d["tools"].items():
        t = tgt["tools"].setdefault(tool, {"calls": 0, "ctx": 0})
        t["calls"] += v["calls"]
        t["ctx"] += v["ctx"]
    r = tgt.setdefault("reads", blank_reads())
    src = d.get("reads") or blank_reads()
    for cls in ("whole", "section"):
        for k, v in src[cls].items():
            r[cls][k] = r[cls].get(k, 0) + v
    r["grep"] += src.get("grep", 0)


def merge_days(all_days, days):
    for day, d in days.items():
        merge_bucket(all_days.setdefault(day, {"models": {}, "tools": {}, "reads": blank_reads()}), d)


def cost_usd(model, m):
    p = PRICING.get(model)
    if not p:
        return None
    return (m["in"] * p[0] + m["out"] * p[1] + m["read"] * p[2] * p[0]
            + m["c5"] * 1.25 * p[0] + m["c1"] * 2.0 * p[0]) / 1e6


def period_keys(day):
    dt = datetime.date.fromisoformat(day)
    iso = dt.isocalendar()
    return [("day", day), ("week", "%d-W%02d" % (iso[0], iso[1])),
            ("month", day[:7]), ("all", "all")]


def build_rows(all_days, ws):
    """day rows -> aggregated CSV rows for every period granularity."""
    agg = {}   # (ptype, period) -> {"models": {}, "tools": {}}
    for day, d in all_days.items():
        for ptype, period in period_keys(day):
            merge_bucket(agg.setdefault(
                (ptype, period), {"models": {}, "tools": {}, "reads": blank_reads()}), d)
    rows = []
    order = {"day": 0, "week": 1, "month": 2, "all": 3}
    for (ptype, period) in sorted(agg, key=lambda k: (order[k[0]], k[1])):
        d = agg[(ptype, period)]
        total = blank_model()
        total_cost, cost_known = 0.0, True
        total_w = 0.0
        for key in sorted(d["models"]):
            model, role = key.split("|")
            m = d["models"][key]
            c = cost_usd(model, m)
            if c is None:
                cost_known = False
            else:
                total_cost += c
            for k in total:
                total[k] += m[k]
            w = weighted(model, m)
            total_w += w
            tot = m["in"] + m["out"] + m["read"] + m["c5"] + m["c1"]
            rows.append([ptype, period, ws, "model", model, role, m["count"],
                         m["in"], m["out"], m["think"], m["read"],
                         m["c5"] + m["c1"], "",
                         "%.2f" % c if c is not None else "",
                         tot, tot // (m["count"] or 1), int(w),
                         m["miss_n"], m["miss_tok"]])
            if ptype == "all":
                MAXP[(ws, model, role)] = m["maxp"]
            elif ptype == "month":
                MAXP_M[(ws, period, model, role)] = m["maxp"]
        if ptype in ("month", "all"):
            # per-request averages: count = msgs, the token columns are
            # per message, context_est_tok = the average box, est$ per msg
            for key in sorted(d["models"]):
                model, role = key.split("|")
                m = d["models"][key]
                n = m["count"] or 1
                c = cost_usd(model, m)
                rows.append([ptype, period, ws, "avg_req", model, role,
                             m["count"], m["in"] // n, m["out"] // n,
                             m["think"] // n, m["read"] // n,
                             (m["c5"] + m["c1"]) // n, box(m) // n,
                             "%.4f" % (c / n) if c is not None else "",
                             "", "", int(weighted(model, m) // n), "", ""])
        tool_ctx = 0
        for tool in sorted(d["tools"], key=lambda t: -d["tools"][t]["ctx"]):
            v = d["tools"][tool]
            tool_ctx += v["ctx"]
            rows.append([ptype, period, ws, "tool", tool, "-", v["calls"],
                         "", "", "", "", "", v["ctx"], "",
                         v["ctx"], v["ctx"] // (v["calls"] or 1), "", "", ""])
        # THE TOTAL ROW sits LAST in its period (the CEO 2026-09-10: "in
        # line under their final record"); the .xlsx bolds it and rules a
        # thick border under it.
        tot = total["in"] + total["out"] + total["read"] + total["c5"] + total["c1"]
        rows.append([ptype, period, ws, "total", "all", "all", total["count"],
                     total["in"], total["out"], total["think"], total["read"],
                     total["c5"] + total["c1"], tool_ctx,
                     "%.2f" % total_cost if cost_known else "",
                     tot, tot // (total["count"] or 1), int(total_w),
                     total["miss_n"], total["miss_tok"]])
    return rows


def keep_other_ws(ws, path=CSV_OUT, fields=FIELDS, ws_col=2):
    """Rows from the other workstation survive a rerun here untouched."""
    if not os.path.isfile(path):
        return []
    with open(path, encoding="utf-8", newline="") as fh:
        return [r for r in list(csv.reader(fh))[1:]
                if len(r) == len(fields) and r[ws_col] != ws]


def local_clock(ts):
    try:
        dt = datetime.datetime.fromisoformat(ts.replace("Z", "+00:00"))
        return dt.astimezone().strftime("%H:%M")
    except (ValueError, AttributeError):
        return ""


def minutes_between(a, b):
    try:
        ta = datetime.datetime.fromisoformat(a.replace("Z", "+00:00"))
        tb = datetime.datetime.fromisoformat(b.replace("Z", "+00:00"))
        return max(0.0, (tb - ta).total_seconds() / 60.0)
    except (ValueError, AttributeError):
        return 0.0


def employee_rows(cache, ws):
    """One row per sub-agent transcript (a file under <session>/subagents/
    IS one employee run). Sorted newest first."""
    rows = []
    for path, state in cache["files"].items():
        parts = path.replace("\\", "/").split("/")
        if "subagents" not in parts or not state.get("days"):
            continue
        parent = parts[parts.index("subagents") - 1]
        total = blank_model()
        models = {}
        calls = ctx = 0
        for d in state["days"].values():
            for key, m in d["models"].items():
                models[key.split("|")[0]] = models.get(key.split("|")[0], 0) + m["count"]
                for k in total:
                    total[k] = max(total[k], m[k]) if k == "maxp" else total[k] + m[k]
            for v in d["tools"].values():
                calls += v["calls"]
                ctx += v["ctx"]
        if not total["count"]:
            continue
        model = max(models, key=models.get)
        cost = cost_usd(model, total)
        day = local_day(state.get("t_first") or "") or sorted(state["days"])[0]
        started = local_clock(state.get("t_first") or "")
        rows.append([ws, day, started,
                     "%.1f" % minutes_between(state.get("t_first", ""), state.get("t_last", "")),
                     parent[:8], model, total["count"], total["in"], total["out"],
                     total["think"], total["read"], total["c5"] + total["c1"],
                     calls, ctx, box(total) // (total["count"] or 1),
                     "%.2f" % cost if cost is not None else "",
                     state.get("label") or ""])
    rows.sort(key=lambda r: (r[1], r[2]), reverse=True)
    return rows


def ktok(n):
    return "~%dk" % round(int(n) / 1000.0) if int(n) >= 1000 else str(n)


def avg_tag(n):
    """Like ktok but always tilde-prefixed, for a per-call average figure."""
    n = int(n)
    return "~%dk" % round(n / 1000.0) if n >= 1000 else "~%d" % n


def write_txt(rows, emp_rows, ws):
    by = {}
    for r in rows:
        by.setdefault((r[0], r[1], r[2]), []).append(r)
    lines = [
        "THE USAGE SHEET - harness-metered token usage, aggregated",
        "regenerated %s by tools/usage_report.py (CSV twin: usage_metrics.csv)"
        % datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "",
        "READING IT: WEIGHTED is the headline - the tokens that count against",
        "the plan (input x1, cache write x1.25 or x2 on the 1h TTL, cache read",
        "x0.1 or x0.025 on Fable, output x5); every other column is raw trivia",
        "beside it. misses = prefix rewrites after a session's first request",
        "(the cache expired or an early byte changed) and the write tokens they",
        "cost. A tool's cost is the context its results inject (ctx_est,",
        "chars/4) times every later re-read. est$ prices known models only;",
        "the billing dashboard is the only dollar truth.",
        "",
        "SECTIONS: totals (all-time / month / week / day), then THE",
        "BREAKDOWNS at the bottom - the weighted pillars, cache misses, the",
        "read diet, averages per request, the context window by month,",
        "averages per tool call, employee runs, averages per day. THE DAILY",
        "LINE with day-vs-previous-days verdicts is usage_daily.txt.",
        ""]
    order = {"all": 0, "month": 1, "week": 2, "day": 3}
    titles = {"all": "ALL-TIME", "month": "BY MONTH", "week": "BY WEEK",
              "day": "BY DAY (newest last)"}
    last_ptype = None
    day_count = {}
    for (ptype, period, ws) in sorted(
            by, key=lambda k: (order[k[0]], k[1], k[2])):
        if ptype != last_ptype:
            lines += ["=" * 66, titles[ptype], "=" * 66]
            last_ptype = ptype
        lines.append("-- %s %s" % (period, ws))
        lines.append("   %-8s %-6s %6s %10s %9s %9s %8s %10s %10s %6s %8s" % (
            "model", "role", "msgs", "WEIGHTED", "input", "output", "think",
            "cache_rd", "cache_wr", "misses", "est$"))
        for r in by[(ptype, period, ws)]:
            if r[3] == "model":
                lines.append("   %-8s %-6s %6s %10s %9s %9s %8s %10s %10s %6s %8s" % (
                    r[4], r[5], r[6], ktok(r[16]), ktok(r[7]), ktok(r[8]), ktok(r[9]),
                    ktok(r[10]), ktok(r[11]), r[17] or "0", r[13] or "-"))
        for r in by[(ptype, period, ws)]:
            if r[3] == "total":
                lines.append("   %-8s %-6s %6s %10s %9s %9s %8s %10s %10s %6s %8s" % (
                    "TOTAL", "", r[6], ktok(r[16]), ktok(r[7]), ktok(r[8]), ktok(r[9]),
                    ktok(r[10]), ktok(r[11]), r[17] or "0", r[13] or "-"))
                if ptype == "day":
                    day_count[ws] = day_count.get(ws, 0) + 1
        tools = [r for r in by[(ptype, period, ws)] if r[3] == "tool"]
        if tools and ptype in ("all", "month"):
            lines.append("   tools by context injected: " + ", ".join(
                "%s x%s %s" % (r[4], r[6], ktok(r[12])) for r in tools[:10]))
        elif tools:
            lines.append("   top tools: " + ", ".join(
                "%s x%s %s" % (r[4], r[6], ktok(r[12])) for r in tools[:5]))
        lines.append("")
    lines += pillar_lines(by)
    lines += diet_lines(ALL_DAYS, ws)
    lines += breakdown_lines(by, emp_rows)
    lines += ["=" * 66, "AVERAGES (per active day)", "=" * 66]
    for (ptype, period, ws) in sorted(by):
        if ptype != "all":
            continue
        n = day_count.get(ws) or 1
        for r in by[(ptype, period, ws)]:
            if r[3] == "total":
                lines.append(
                    "%s: %d active days | avg/day: %s WEIGHTED | %s msgs, %s in, %s out,"
                    " %s cache_rd, %s cache_wr, %.1f misses" % (
                        ws, n, ktok(int(r[16]) // n), int(r[6]) // n, ktok(int(r[7]) // n),
                        ktok(int(r[8]) // n), ktok(int(r[10]) // n),
                        ktok(int(r[11]) // n), int(r[17] or 0) / float(n)))
    with open(TXT_OUT, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")


def pillars(model, m):
    """The four weighted pillars of a bucket -> {name: weighted tokens}."""
    return {"cache writes": m["c5"] * W_WRITE_5M + m["c1"] * W_WRITE_1H,
            "cache reads": m["read"] * read_mult(model),
            "output": m["out"] * W_OUTPUT,
            "input": m["in"] * W_INPUT}


def pillar_lines(by):
    """THE WEIGHTED PILLARS + CACHE MISSES (ideas 13 + 14), all-time and by
    month, from the CSV rows (weighted per model row; pillars need the
    raw columns, so they are recomputed here)."""
    L = []
    H = "=" * 66
    L += [H, "THE WEIGHTED PILLARS (what actually hits the budget, by month)", H,
          "   misses = prefix rewrites after the first request; their write tokens are the waste nobody chose"]
    for (ptype, period, ws) in sorted(by, key=lambda k: (k[2], k[0] != "all", k[1])):
        if ptype not in ("all", "month"):
            continue
        agg = {"cache writes": 0.0, "cache reads": 0.0, "output": 0.0, "input": 0.0}
        raw = w = 0
        miss_n = miss_tok = 0
        for r in by[(ptype, period, ws)]:
            if r[3] != "model":
                continue
            m = {"in": int(r[7]), "out": int(r[8]), "read": int(r[10]),
                 "c5": int(r[11]), "c1": 0}
            for k, v in pillars(r[4], m).items():
                agg[k] += v
            raw += int(r[14])
            w += int(r[16])
            miss_n += int(r[17] or 0)
            miss_tok += int(r[18] or 0)
        if not w:
            continue
        top = sorted(agg.items(), key=lambda kv: -kv[1])
        L.append("-- %s %s: %s weighted (raw %s, %.0f%%) | %s | misses %d (%s written, %s weighted)" % (
            period if ptype == "month" else "ALL-TIME", ws, ktok(w), ktok(raw),
            100.0 * w / (raw or 1),
            ", ".join("%s %.0f%%" % (k, 100.0 * v / w) for k, v in top),
            miss_n, ktok(miss_tok), ktok(int(miss_tok * W_WRITE_5M))))
    L.append("   (pillars approximate the 5-minute write rate; a 1-hour-TTL session writes at 2x)")
    L.append("")
    return L


def read_diet(d):
    """One day's reads -> (whole_n, whole_ctx, big_n, section_n, section_ctx, grep_n)."""
    r = d.get("reads") or blank_reads()
    return (r["whole"]["n"], r["whole"]["ctx"], r["whole"].get("big", 0),
            r["section"]["n"], r["section"]["ctx"], r.get("grep", 0))


def day_summary(day, d):
    """The numbers the daily line judges: weighted, raw, pillars, misses, reads."""
    agg = {"cache writes": 0.0, "cache reads": 0.0, "output": 0.0, "input": 0.0}
    raw = w = miss_n = miss_tok = msgs = 0
    for key, m in d["models"].items():
        model = key.split("|")[0]
        for k, v in pillars(model, m).items():
            agg[k] += v
        w += weighted(model, m)
        raw += m["in"] + m["out"] + m["read"] + m["c5"] + m["c1"]
        miss_n += m["miss_n"]
        miss_tok += m["miss_tok"]
        msgs += m["count"]
    wn, wctx, big, sn, sctx, grep = read_diet(d)
    read_n = wn + sn
    return {"day": day, "weighted": w, "raw": raw, "pillars": agg, "msgs": msgs,
            "miss_n": miss_n, "miss_tok": miss_tok, "whole_n": wn,
            "whole_ctx": wctx, "big_n": big, "section_n": sn,
            "section_ctx": sctx, "grep_n": grep,
            "section_share": (100.0 * sn / (sn + wn)) if (sn + wn) else None,
            "whole_avg": (wctx // wn) if wn else 0,
            "read_avg": ((wctx + sctx) // read_n) if read_n else None}


def _verdict(value, baseline, higher_is_bad=True):
    """-> (label, pct) comparing a day's value to the prior-days average."""
    if not baseline:
        return ("n/a", 0)
    ratio = value / float(baseline)
    pct = int(round((ratio - 1.0) * 100))
    if ratio >= HIGH_RATIO:
        return ("HIGH" if higher_is_bad else "GOOD", pct)
    if ratio <= LOW_RATIO:
        return ("LOW" if higher_is_bad else "POOR", pct)
    return ("normal", pct)


def daily_lines(all_days, ws):
    """THE DAILY LINE (the CEO 2026-09-10: "a number without comparison means
    nothing"): every active day judged against the previous COMPARE_DAYS
    active days - weighted total, cache misses, whole-file reads - with a
    HIGH / normal / LOW verdict and the reason. Oldest first; the tail is
    what standup shows."""
    days = sorted(all_days)
    sums = [day_summary(d, all_days[d]) for d in days]
    today = datetime.date.today().isoformat()
    L = ["# THE DAILY LINE - one line per active day, judged against the previous "
         "%d active days (regenerated by tools/usage_report.py; read the TAIL)" % COMPARE_DAYS,
         "# date | WS | weighted (raw) | vs prev days | top pillar | misses | reads: whole (big) / section, share, avg tokens/read | verdict"]
    for i, s in enumerate(sums):
        prev = sums[max(0, i - COMPARE_DAYS):i]
        base_w = sum(p["weighted"] for p in prev) / len(prev) if prev else 0
        base_miss = sum(p["miss_n"] for p in prev) / float(len(prev)) if prev else 0
        base_big = sum(p["big_n"] for p in prev) / float(len(prev)) if prev else 0
        base_share = [p["section_share"] for p in prev if p["section_share"] is not None]
        base_share = sum(base_share) / len(base_share) if base_share else None
        wl, wp = _verdict(s["weighted"], base_w)
        top = max(s["pillars"].items(), key=lambda kv: kv[1])
        top_txt = "%s %.0f%%" % (top[0], 100.0 * top[1] / (s["weighted"] or 1))
        reasons = []
        if wl == "HIGH":
            reasons.append("spend +%d%% vs prev %d days" % (wp, len(prev)))
        elif wl == "LOW":
            reasons.append("spend %d%% vs prev %d days (good)" % (wp, len(prev)))
        if prev and s["miss_n"] >= max(3, 2 * base_miss):
            reasons.append("cache misses %d vs avg %.1f (%s wasted)" % (
                s["miss_n"], base_miss, ktok(int(s["miss_tok"] * W_WRITE_5M))))
        if prev and s["big_n"] >= max(3, 2 * base_big):
            reasons.append("heavy whole-file reads %d vs avg %.1f" % (s["big_n"], base_big))
        if base_share is not None and s["section_share"] is not None \
                and s["section_share"] < base_share - 20:
            reasons.append("section-read share %.0f%% vs avg %.0f%%" % (s["section_share"], base_share))
        if not prev:
            verdict = "first day - no baseline"
        elif wl == "HIGH" or (reasons and wl != "LOW"):
            verdict = "CHECK: " + "; ".join(reasons)
        elif wl == "LOW":
            verdict = "LOW spend - " + "; ".join(reasons)
        else:
            verdict = "normal"
        partial = " (today, partial)" if s["day"] == today else ""
        vs = ("%+d%%" % wp) if prev else "-"
        reads_txt = ("%.0f%% sectioned avg %s" % (s["section_share"], avg_tag(s["read_avg"])) ) \
            if s["section_share"] is not None else "no reads"
        L.append("%s | %s | %s (%s) | %s | %s | %d | %d (%d) / %d, %s | %s%s" % (
            s["day"], ws, ktok(int(s["weighted"])), ktok(s["raw"]), vs, top_txt,
            s["miss_n"], s["whole_n"], s["big_n"], s["section_n"],
            reads_txt,
            verdict, partial))
    return L


def diet_lines(all_days, ws):
    """THE READ DIET by month (idea 17): section vs whole-file reads and the
    context each injected; the daily comparison lives in usage_daily.txt."""
    L = []
    H = "=" * 66
    L += [H, "THE READ DIET BY MONTH (section reads vs whole-file reads; big = a whole read past %s tokens)" % ktok(BIG_READ_TOK), H,
          "   %-8s %8s %9s %6s %10s %8s %9s %8s %8s" % (
              "month", "whole", "whole_ctx", "big", "whole_avg", "section", "sect_ctx", "greps", "share")]
    months = {}
    for day, d in all_days.items():
        mo = months.setdefault(day[:7], blank_reads())
        src = d.get("reads") or blank_reads()
        for cls in ("whole", "section"):
            for k, v in src[cls].items():
                mo[cls][k] = mo[cls].get(k, 0) + v
        mo["grep"] += src.get("grep", 0)
    for mo in sorted(months):
        r = months[mo]
        wn, sn = r["whole"]["n"], r["section"]["n"]
        L.append("   %-8s %8d %9s %6d %10s %8d %9s %8d %7s" % (
            mo, wn, ktok(r["whole"]["ctx"]), r["whole"].get("big", 0),
            ktok(r["whole"]["ctx"] // wn) if wn else "-", sn, ktok(r["section"]["ctx"]),
            r["grep"], ("%.0f%%" % (100.0 * sn / (sn + wn))) if (sn + wn) else "-"))
    L.append("")
    return L


def breakdown_lines(by, emp_rows):
    """THE BREAKDOWNS (the CEO's ask 2026-09-10): per request, the context
    window by month, per tool call, employee runs."""
    L = []
    H = "=" * 66
    # 1. per request
    L += [H, "AVERAGES PER REQUEST (all-time; one request = one API call)", H,
          "   box = the prompt the model read that turn (input + cache_rd + cache_wr)"]
    for (ptype, period, ws) in sorted(by):
        if ptype != "all":
            continue
        L.append("-- %s" % ws)
        L.append("   %-8s %-6s %6s %8s %8s %8s %9s %9s %9s %9s %8s" % (
            "model", "role", "msgs", "in/req", "out/req", "think", "cache_rd",
            "cache_wr", "avg_box", "max_box", "est$/req"))
        maxes = {r[4] + "|" + r[5]: r for r in by[(ptype, period, ws)] if r[3] == "model"}
        for r in by[(ptype, period, ws)]:
            if r[3] != "avg_req":
                continue
            L.append("   %-8s %-6s %6s %8s %8s %8s %9s %9s %9s %9s %8s" % (
                r[4], r[5], r[6], ktok(r[7]), ktok(r[8]), ktok(r[9]), ktok(r[10]),
                ktok(r[11]), ktok(r[12]), ktok(MAXP.get((ws, r[4], r[5]), 0)),
                r[13] or "-"))
    L.append("")
    # 2. the context window by month (main role: what the manager re-reads)
    L += [H, "THE CONTEXT WINDOW BY MONTH (avg box per request; a rising avg = sessions that should have checkpointed)", H]
    for (ptype, period, ws) in sorted(by):
        if ptype != "month":
            continue
        cells = []
        for r in by[(ptype, period, ws)]:
            if r[3] == "avg_req":
                cells.append("%s/%s avg %s max %s (%s req)" % (
                    r[4], r[5], ktok(r[12]), ktok(MAXP_M.get((ws, period, r[4], r[5]), 0)), r[6]))
        L.append("-- %s %s: %s" % (period, ws, "; ".join(cells) if cells else "-"))
    L.append("")
    # 3. per tool call
    L += [H, "AVERAGES PER TOOL CALL (all-time; ctx = what the result injected, chars/4)", H]
    for (ptype, period, ws) in sorted(by):
        if ptype != "all":
            continue
        tools = [r for r in by[(ptype, period, ws)] if r[3] == "tool"]
        total_ctx = sum(int(r[12] or 0) for r in tools) or 1
        L.append("-- %s" % ws)
        L.append("   %-22s %7s %10s %10s %6s" % ("tool", "calls", "ctx_total", "ctx/call", "share"))
        for r in sorted(tools, key=lambda r: -int(r[12] or 0)):
            calls = int(r[6] or 0) or 1
            L.append("   %-22s %7s %10s %10s %5.1f%%" % (
                r[4][:22], r[6], ktok(r[12]), ktok(int(r[12] or 0) // calls),
                100.0 * int(r[12] or 0) / total_ctx))
    L.append("")
    # 4. employee runs
    L += [H, "EMPLOYEE RUNS (one sub-agent transcript = one run; full list: usage_employees.csv)", H]
    by_ws = {}
    for r in emp_rows:
        by_ws.setdefault(r[0], []).append(r)
    for ws in sorted(by_ws):
        runs = by_ws[ws]
        n = len(runs)
        L.append("-- %s: %d runs" % (ws, n))
        by_model = {}
        for r in runs:
            by_model.setdefault(r[5], []).append(r)
        L.append("   %-8s %5s %7s %8s %8s %9s %9s %7s %8s %8s" % (
            "model", "runs", "avg_min", "avg_msgs", "avg_in", "avg_out", "avg_cache", "avg_tool", "avg_box", "avg_est$"))
        for model in sorted(by_model):
            rs = by_model[model]
            k = len(rs)
            costs = [float(r[15]) for r in rs if r[15]]
            L.append("   %-8s %5d %7.1f %8d %8s %9s %9s %7d %8s %8s" % (
                model, k, sum(float(r[3]) for r in rs) / k,
                sum(int(r[6]) for r in rs) // k, ktok(sum(int(r[7]) for r in rs) // k),
                ktok(sum(int(r[8]) for r in rs) // k), ktok(sum(int(r[10]) for r in rs) // k),
                sum(int(r[12]) for r in rs) // k, ktok(sum(int(r[14]) for r in rs) // k),
                "%.2f" % (sum(costs) / len(costs)) if costs else "-"))
        L.append("   costliest runs:")
        for r in sorted(runs, key=lambda r: -(float(r[15]) if r[15] else 0))[:10]:
            L.append("   %s %s %-7s %5s min %4s tools %7s box  $%-6s %s" % (
                r[1], r[2], r[5], r[3], r[12], ktok(r[14]), r[15] or "-", r[16][:60]))
    if not by_ws:
        L.append("   (no sub-agent transcripts found)")
    L.append("")
    return L


ALL_DAYS = {}  # day -> bucket, set by main() for the diet section
MAXP = {}     # (ws, model, role) -> biggest prompt all-time
MAXP_M = {}   # (ws, month, model, role) -> biggest prompt that month


XLSX_NUMBER = "#,##0"        # the Format Cells screenshot: Number, 1000 separator
XLSX_MONEY = "#,##0.00"
NUMERIC_COLS = {"count", "input_tok", "output_tok", "think_tok",
                "cache_read_tok", "cache_create_tok", "context_est_tok",
                "total_tok", "avg_tok", "msgs", "tool_calls", "ctx_est_tok",
                "avg_box_tok", "weighted_tok", "cache_miss_n", "cache_miss_tok"}
MONEY_COLS = {"est_cost_usd"}


def _xlsx_sheet(wb, title, fields, rows, total_scope_col=None, block_key=None):
    """THE SPREADSHEET RULE (REPORTING_METHOD.md): header frozen, numbers
    formatted with separators, TOTAL rows bold + thick bottom border,
    one thick border after every period block."""
    from openpyxl.styles import Font, Border, Side, Alignment
    ws_ = wb.create_sheet(title)
    ws_.append(fields)
    bold = Font(bold=True)
    thick = Border(bottom=Side(style="thick"))
    for c in ws_[1]:
        c.font = bold
        c.alignment = Alignment(horizontal="center")
    ws_.freeze_panes = "A2"
    ws_.auto_filter.ref = "A1:%s1" % ws_.cell(row=1, column=len(fields)).column_letter
    fmt = {}
    for i, f in enumerate(fields, 1):
        if f in NUMERIC_COLS:
            fmt[i] = XLSX_NUMBER
        elif f in MONEY_COLS:
            fmt[i] = XLSX_MONEY
    for r in rows:
        vals = []
        for i, v in enumerate(r, 1):
            if i in fmt and v not in ("", None):
                try:
                    v = float(v) if i in fmt and fmt[i] == XLSX_MONEY else int(float(v))
                except (TypeError, ValueError):
                    pass
            vals.append(v)
        ws_.append(vals)
        row = ws_.max_row
        for i in fmt:
            ws_.cell(row=row, column=i).number_format = fmt[i]
        if total_scope_col is not None and r[total_scope_col] == "total":
            for c in ws_[row]:
                c.font = bold
                c.border = thick
    widths = {"period_type": 11, "period": 12, "ws": 5, "scope": 8, "name": 16,
              "role": 6, "brief": 60, "parent_session": 14, "started": 8}
    for i, f in enumerate(fields, 1):
        ws_.column_dimensions[ws_.cell(row=1, column=i).column_letter].width = widths.get(f, 13)
    return ws_


def write_xlsx(rows, emp_rows):
    """usage_metrics.xlsx: sheets usage (model + tool rows, TOTAL last per
    period), per_request (the avg_req rows), employees. Returns a note."""
    try:
        import openpyxl
    except ImportError:
        return "openpyxl missing - no .xlsx (python -m pip install openpyxl)"
    order = {"day": 0, "week": 1, "month": 2, "all": 3}
    scope_rank = {"model": 0, "tool": 1, "total": 2}
    main = sorted([r for r in rows if r[3] in scope_rank],
                  key=lambda r: (order.get(r[0], 9), r[1], r[2], scope_rank[r[3]]))
    per_req = sorted([r for r in rows if r[3] == "avg_req"],
                     key=lambda r: (order.get(r[0], 9), r[1], r[2], r[4], r[5]))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _xlsx_sheet(wb, "usage", FIELDS, main, total_scope_col=3)
    _xlsx_sheet(wb, "per_request", FIELDS, per_req)
    _xlsx_sheet(wb, "employees", EMP_FIELDS, emp_rows)
    try:
        wb.save(XLSX_OUT)
    except PermissionError:
        return ("%s NOT rewritten - it is open in Excel; close it and rerun "
                "(CSV/TXT are current)" % os.path.relpath(XLSX_OUT, ROOT))
    return os.path.relpath(XLSX_OUT, ROOT)


def parse_ts(ts):
    """Transcript ISO timestamp ('...Z') -> aware datetime, or None."""
    try:
        return datetime.datetime.fromisoformat((ts or "").replace("Z", "+00:00"))
    except (ValueError, AttributeError):
        return None


def parse_git_ts(s):
    """git log '%ci' ('YYYY-MM-DD HH:MM:SS +ZZZZ') -> aware datetime, or None."""
    try:
        dt = datetime.datetime.strptime(s[:19], "%Y-%m-%d %H:%M:%S")
        sign = 1 if s[20] == "+" else -1
        oh, om = int(s[21:23]), int(s[23:25])
        return dt.replace(tzinfo=datetime.timezone(
            sign * datetime.timedelta(hours=oh, minutes=om)))
    except (ValueError, IndexError):
        return None


def git_commits():
    """All commits on HEAD, oldest first: (short_hash, dt, subject). Powers
    THE PER-ARC LINE - an arc is the span between two consecutive
    "Checkpoint:" commits (the /checkpoint skill's own commit convention)."""
    try:
        out = subprocess.check_output(
            ["git", "log", "--format=%h|%ci|%s"], cwd=ROOT,
            encoding="utf-8", errors="replace")
    except (OSError, subprocess.CalledProcessError):
        return []
    commits = []
    for line in out.splitlines():
        parts = line.split("|", 2)
        if len(parts) != 3:
            continue
        dt = parse_git_ts(parts[1])
        if dt:
            commits.append((parts[0], dt, parts[2]))
    commits.reverse()
    return commits


def keep_other_ws_lines(path, ws, ws_field_idx=1):
    """Non-comment lines from another workstation survive a rerun here -
    the plain-text twin of keep_other_ws, for '|'-delimited ledgers that
    are regenerated whole (not appended) each run."""
    if not os.path.isfile(path):
        return []
    kept = []
    with open(path, encoding="utf-8") as fh:
        for raw in fh:
            ln = raw.rstrip("\n")
            if not ln or ln.startswith("#"):
                continue
            parts = [p.strip() for p in ln.split("|")]
            if len(parts) > ws_field_idx and parts[ws_field_idx] != ws:
                kept.append(ln)
    return kept


def arc_lines(all_reqs, emp_starts, ws):
    """THE PER-ARC LINE (the CEO 2026-09-14: "Per-arc cost line, sounds
    great"): one line per arc between consecutive "Checkpoint:" commits -
    commits in the window, weighted+raw spend, employee runs started in
    the window, spend per commit, the biggest weighted pillar, and a CHECK
    verdict when per-commit spend beats the mean of the previous
    BASELINE_ARCS arcs by more than (ARC_OVERSPEND - 1) x 100%. The open
    arc since the last checkpoint is the final line, marked (open)."""
    commits = git_commits()
    if not commits:
        return []
    checkpoints = [c for c in commits if c[2].startswith("Checkpoint:")]
    reqs = []
    for (ts, model, _role, i_, o_, rd, c5, c1) in all_reqs:
        dt = parse_ts(ts)
        if dt:
            reqs.append((dt, model, i_, o_, rd, c5, c1))

    def in_window(dt, start_dt, end_dt):
        if start_dt is not None and dt <= start_dt:
            return False
        if end_dt is not None and dt > end_dt:
            return False
        return True

    spans = []  # (start_dt, start_hash, end_commit_or_None)
    for i, cp in enumerate(checkpoints):
        if i > 0:
            start_dt, start_hash = checkpoints[i - 1][1], checkpoints[i - 1][0]
        else:
            start_dt, start_hash = None, commits[0][0]
        spans.append((start_dt, start_hash, cp))
    if checkpoints:
        open_start_dt, open_start_hash = checkpoints[-1][1], checkpoints[-1][0]
    else:
        open_start_dt, open_start_hash = None, commits[0][0]
    spans.append((open_start_dt, open_start_hash, None))

    lines = []
    closed_per_commit = []
    for start_dt, start_hash, end in spans:
        is_open = end is None
        end_dt = end[1] if not is_open else None
        end_hash = "HEAD" if is_open else end[0]
        date = datetime.date.today().isoformat() if is_open else end[1].date().isoformat()
        cms = [c for c in commits if in_window(c[1], start_dt, end_dt)]
        n_commits = len(cms)
        agg = {"cache writes": 0.0, "cache reads": 0.0, "output": 0.0, "input": 0.0}
        w = raw = 0.0
        for dt, model, i_, o_, rd, c5, c1 in reqs:
            if not in_window(dt, start_dt, end_dt):
                continue
            m = {"in": i_, "out": o_, "read": rd, "c5": c5, "c1": c1}
            w += weighted(model, m)
            raw += i_ + o_ + rd + c5 + c1
            for k, v in pillars(model, m).items():
                agg[k] += v
        emps = sum(1 for dt in emp_starts if in_window(dt, start_dt, end_dt))
        per_commit = w / max(n_commits, 1)
        top = max(agg.items(), key=lambda kv: kv[1]) if w else None
        top_txt = ("%s %.0f%%" % (top[0], 100.0 * top[1] / (w or 1))) if top else "-"
        base = closed_per_commit[-BASELINE_ARCS:]
        if not base:
            verdict = "first arcs - no baseline"
        else:
            base_mean = sum(base) / len(base)
            if base_mean and per_commit > base_mean * ARC_OVERSPEND:
                pct = int(round((per_commit / base_mean - 1.0) * 100))
                verdict = "CHECK: per-commit +%d%% vs previous %d arcs" % (pct, len(base))
            else:
                verdict = "normal"
        if not is_open:
            closed_per_commit.append(per_commit)
        lines.append("%s | %s | %s..%s | %d | %s (%s) | %d | %s | %s | %s%s" % (
            date, ws, start_hash, end_hash, n_commits,
            ktok(int(w)), ktok(int(raw)), emps, ktok(int(per_commit)), top_txt,
            verdict, " (open)" if is_open else ""))
    return lines


def miss_lines(all_misses, ws):
    """THE CACHE-MISS INVESTIGATION (the CEO 2026-09-14, "come up with a
    plan, create data points"): one line per detected cache miss - same
    detection as the daily line's miss/cold counters (MISS_MIN_BOX /
    MISS_SHARE, unchanged), itemized with a likely cause so the dominant
    waste is visible without a script. Regenerated whole each run from
    this workstation's own cached transcripts (each machine only has its
    own transcripts to mine)."""
    rows = []
    for (ts, sess, gap_min, bx, wr, wasted, cause) in all_misses:
        dt = parse_ts(ts)
        if not dt:
            continue
        local = dt.astimezone()
        gap_txt = "-" if gap_min is None else "%.1f" % gap_min
        rows.append((local.strftime("%Y-%m-%d %H:%M"), sess, gap_txt, bx, wr, wasted, cause))
    rows.sort(key=lambda r: r[0])
    total_n, total_wasted = len(rows), sum(r[5] for r in rows)
    by_cause = {}
    for r in rows:
        c = by_cause.setdefault(r[6], {"n": 0, "wasted": 0.0})
        c["n"] += 1
        c["wasted"] += r[5]
    L = ["# THE CACHE-MISS INVESTIGATION - one line per detected cache miss "
         "(detection unchanged: MISS_MIN_BOX/MISS_SHARE in tools/usage_report.py; "
         "this ledger itemizes and classifies each one; regenerated whole each run "
         "from this workstation's own transcripts)",
         "# date time | ws | session (8 chars) | gap since previous request (min) | "
         "box tokens | cache write tokens | wasted ~weighted | likely cause",
         "# SUMMARY: %d misses, %s wasted ~weighted, split by cause:" % (
             total_n, ktok(int(total_wasted)))]
    for cause in sorted(by_cause, key=lambda c: -by_cause[c]["wasted"]):
        c = by_cause[cause]
        L.append("#   %s: %d misses, %s wasted" % (cause, c["n"], ktok(int(c["wasted"]))))
    for r in rows:
        L.append("%s | %s | %s | %s | %d | %d | %s | %s" % (
            r[0], ws, r[1], r[2], r[3], r[4], ktok(int(r[5])), r[6]))
    return L, total_n, total_wasted, by_cause


def main():
    quiet = "--quiet" in sys.argv[1:]
    ws = workstation()
    cache = {"version": CACHE_VERSION, "files": {}}
    if os.path.isfile(CACHE):
        try:
            loaded = json.load(open(CACHE, encoding="utf-8"))
            if loaded.get("version") == CACHE_VERSION:
                cache = loaded
        except ValueError:
            pass
    all_days = {}
    n_files = 0
    for tdir in transcript_dirs():
        for base, _dirs, names in os.walk(tdir):
            for name in sorted(names):
                if not name.endswith(".jsonl"):
                    continue
                path = os.path.join(base, name)
                state = cache["files"].setdefault(path, {})
                parse_file(path, state,
                           force_agent="subagents" in base.lower())
                merge_days(all_days, state["days"])
                n_files += 1
    if not all_days:
        print("No transcripts found under ~/.claude/projects - nothing to do.")
        return
    os.makedirs(HIST, exist_ok=True)
    json.dump(cache, open(CACHE, "w", encoding="utf-8"))
    ALL_DAYS.update(all_days)
    rows = build_rows(all_days, ws) + keep_other_ws(ws)
    with open(CSV_OUT, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(FIELDS)
        w.writerows(rows)
    emp = employee_rows(cache, ws) + keep_other_ws(ws, EMP_OUT, EMP_FIELDS, 0)
    with open(EMP_OUT, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(EMP_FIELDS)
        w.writerows(emp)
    write_txt(rows, emp, ws)
    daily = daily_lines(all_days, ws)
    with open(DAILY_OUT, "w", encoding="utf-8") as fh:
        fh.write("\n".join(daily) + "\n")

    # THE PER-ARC LINE (usage_by_arc.txt): per-request + employee-start
    # timestamps flattened across every cached transcript on this machine.
    all_reqs, emp_starts = [], []
    for fpath, state in cache["files"].items():
        all_reqs.extend(state.get("reqs") or [])
        parts = fpath.replace("\\", "/").split("/")
        if "subagents" in parts and state.get("t_first"):
            dt = parse_ts(state["t_first"])
            if dt:
                emp_starts.append(dt)
    arc_header = [
        "# THE PER-ARC LINE - one line per arc between consecutive "
        "\"Checkpoint:\" commits (commits = count of commits strictly after "
        "the previous checkpoint through and including this arc's own "
        "checkpoint commit; regenerated by tools/usage_report.py)",
        "# date | ws | arc (start..end commit) | commits | weighted (raw) | "
        "employee runs | per commit | top pillar | verdict"]
    arc_body = arc_lines(all_reqs, emp_starts, ws) + keep_other_ws_lines(ARC_OUT, ws)
    arc_body.sort(key=lambda ln: ln.split("|")[0].strip())
    with open(ARC_OUT, "w", encoding="utf-8") as fh:
        fh.write("\n".join(arc_header + arc_body) + "\n")

    # THE CACHE-MISS INVESTIGATION (cache_misses.txt + cache_miss_runs.txt).
    all_misses = []
    for state in cache["files"].values():
        all_misses.extend(state.get("misses") or [])
    miss_body, miss_n, miss_wasted, miss_by_cause = miss_lines(all_misses, ws)
    with open(MISS_OUT, "w", encoding="utf-8") as fh:
        fh.write("\n".join(miss_body) + "\n")
    dominant = max(miss_by_cause, key=lambda c: miss_by_cause[c]["wasted"]) \
        if miss_by_cause else "-"
    run_header = "# date time | ws | misses | wasted ~weighted | dominant cause"
    run_line = "%s | %s | %d | %s | %s" % (
        datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), ws, miss_n,
        ktok(int(miss_wasted)), dominant)
    need_header = not os.path.isfile(MISS_RUNS_OUT)
    # THE REPEATED RUN (systems audit 2026-09-29, TOKENS T1): two hooks run
    # this report minutes apart and the same misses were ledgered twice,
    # doubling the "wasted" figure the daily CHECK quotes. A line whose
    # fields after the stamp equal the previous line's is the same reading.
    prev = ""
    if not need_header:
        with open(MISS_RUNS_OUT, encoding="utf-8") as fh:
            data = [ln.rstrip() for ln in fh if ln.strip() and not ln.startswith("#")]
        prev = data[-1] if data else ""
    same = prev.split(" | ", 1)[-1] == run_line.split(" | ", 1)[-1]
    with open(MISS_RUNS_OUT, "a", encoding="utf-8") as fh:
        if need_header:
            fh.write(run_header + "\n")
        if not same:
            fh.write(run_line + "\n")

    xlsx_note = write_xlsx(rows, emp)
    if quiet:
        return
    days = sorted(all_days)
    total = blank_model()
    total_w = 0.0
    for d in all_days.values():
        for key, m in d["models"].items():
            total_w += weighted(key.split("|")[0], m)
            for k in total:
                total[k] += m[k]
    print("USAGE SHEET: %d transcripts, %d active days (%s .. %s) on %s" % (
        n_files, len(days), days[0], days[-1], ws))
    print("  WEIGHTED %s (raw %s, %.0f%%) | %s msgs | in %s out %s | cache_rd %s cache_wr %s | misses %d" % (
        ktok(int(total_w)), ktok(total["in"] + total["out"] + total["read"] + total["c5"] + total["c1"]),
        100.0 * total_w / ((total["in"] + total["out"] + total["read"] + total["c5"] + total["c1"]) or 1),
        total["count"], ktok(total["in"]), ktok(total["out"]),
        ktok(total["read"]), ktok(total["c5"] + total["c1"]), total["miss_n"]))
    print("  THE DAILY LINE (last 2 of %s):" % os.path.relpath(DAILY_OUT, ROOT))
    for ln in daily[-2:]:
        print("    " + ln)
    arc_data_lines = [ln for ln in arc_body if not ln.startswith("#")]
    print("  THE PER-ARC LINE (last 3 of %s):" % os.path.relpath(ARC_OUT, ROOT))
    for ln in arc_data_lines[-3:]:
        print("    " + ln)
    print("  CACHE MISSES this run: %d, %s wasted, dominant cause %s (-> %s)" % (
        miss_n, ktok(int(miss_wasted)), dominant, os.path.relpath(MISS_OUT, ROOT)))
    print("  -> OPEN %s (the spreadsheet: usage / per_request / employees)" % xlsx_note)
    print("     %s (totals + the breakdowns, plain text)" % os.path.relpath(TXT_OUT, ROOT))
    print("     csv twins: %s, %s (%d employee runs)" % (
        os.path.relpath(CSV_OUT, ROOT), os.path.relpath(EMP_OUT, ROOT),
        len([r for r in emp if r[0] == ws])))


if __name__ == "__main__":
    main()
